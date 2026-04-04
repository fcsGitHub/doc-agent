"""Document Parse Skill — deterministic DOCX, PDF, and Markdown parsing."""

from __future__ import annotations

import re
import time
from dataclasses import dataclass, field
from typing import Any, Literal

try:
    import pdfplumber  # type: ignore[import-untyped]
except ImportError:  # pragma: no cover — not installed locally
    pdfplumber = None  # type: ignore[assignment]

from skills.base import BaseSkill, SkillContext, SkillResult


@dataclass
class SectionNode:
    """A heading node in the document structure tree."""

    title: str
    level: int
    content: str
    children: list[SectionNode] = field(default_factory=list)


@dataclass
class ParseInput:
    """Input specification for document parsing."""

    file_path: str
    file_type: Literal["docx", "pdf", "md"]


@dataclass
class ParseOutput:
    """Output from document parsing."""

    structure: list[SectionNode]  # top-level sections
    raw_text: str
    metadata: dict[str, Any]  # page_count, word_count, etc.


def _section_node_to_dict(node: SectionNode) -> dict[str, Any]:
    """Recursively convert a SectionNode to a JSON-serializable dict."""
    return {
        "title": node.title,
        "level": node.level,
        "content": node.content,
        "children": [_section_node_to_dict(c) for c in node.children],
    }


class DocumentParseSkill(BaseSkill):
    """Parse DOCX, PDF or Markdown files into structured section trees."""

    name = "document_parse"
    description = "Parse DOCX, PDF or Markdown files into structured section trees"
    prerequisites: list[str] = []

    async def execute(self, context: SkillContext) -> SkillResult:
        """Execute document parsing based on context.input_data."""
        start_ms = int(time.time() * 1000)
        try:
            input_data = context.input_data
            file_path: str = input_data["file_path"]
            file_type: str = input_data["file_type"]

            if file_type == "docx":
                sections, raw_text = self._parse_docx(file_path)
            elif file_type == "pdf":
                sections, raw_text = self._parse_pdf(file_path)
            elif file_type == "md":
                sections, raw_text = self._parse_md(file_path)
            else:
                return SkillResult(
                    success=False,
                    error=f"Unsupported file type: {file_type}",
                    execution_time_ms=int(time.time() * 1000) - start_ms,
                )

            word_count = len(raw_text.split()) if raw_text.strip() else 0
            metadata: dict[str, Any] = {"word_count": word_count}
            if file_type == "pdf":
                metadata["page_count"] = input_data.get("page_count", 0)

            output = ParseOutput(
                structure=sections,
                raw_text=raw_text,
                metadata=metadata,
            )

            # DB persistence — optional
            if context.db_session is not None:
                try:
                    from models.document import ParsedDocument

                    record = ParsedDocument(
                        task_id=context.task_id,
                        source_document_id=input_data.get(
                            "source_document_id", context.task_id
                        ),
                        structure=[_section_node_to_dict(s) for s in sections],
                        raw_text=raw_text,
                        metadata_=metadata,
                    )
                    context.db_session.add(record)
                    await context.db_session.flush()
                except Exception:
                    # Don't fail the skill if DB persistence fails
                    pass

            return SkillResult(
                success=True,
                output={
                    "structure": [_section_node_to_dict(s) for s in sections],
                    "raw_text": raw_text,
                    "metadata": metadata,
                },
                execution_time_ms=int(time.time() * 1000) - start_ms,
            )
        except Exception as exc:
            return SkillResult(
                success=False,
                error=str(exc),
                execution_time_ms=int(time.time() * 1000) - start_ms,
            )

    # ------------------------------------------------------------------
    # DOCX parser
    # ------------------------------------------------------------------
    def _parse_docx(self, file_path: str) -> tuple[list[SectionNode], str]:
        """Parse a DOCX file into section nodes and raw text."""
        from docx import Document

        doc = Document(file_path)
        flat_sections: list[dict[str, Any]] = []
        raw_parts: list[str] = []
        current_content_parts: list[str] = []

        for para in doc.paragraphs:
            style_name = (para.style.name or "") if para.style else ""
            text = para.text.strip()

            # Detect headings: "Heading 1", "Heading 2", etc.
            heading_match = (
                re.match(r"^Heading (\d+)$", style_name) if style_name else None
            )
            if heading_match:
                level = int(heading_match.group(1))
                # Flush accumulated content to previous section
                if flat_sections:
                    flat_sections[-1]["content"] = "\n".join(
                        current_content_parts
                    ).strip()
                current_content_parts = []
                flat_sections.append({"title": text, "level": level, "content": ""})
            else:
                if text:
                    current_content_parts.append(text)

            if text:
                raw_parts.append(text)

        # Flush last section's content
        if flat_sections:
            flat_sections[-1]["content"] = "\n".join(current_content_parts).strip()

        # Handle tables
        for table in doc.tables:
            for row in table.rows:
                row_text = " | ".join(cell.text.strip() for cell in row.cells)
                if row_text.strip():
                    raw_parts.append(row_text)

        raw_text = "\n".join(raw_parts)
        tree = self._build_tree(flat_sections)
        return tree, raw_text

    # ------------------------------------------------------------------
    # PDF parser
    # ------------------------------------------------------------------
    def _parse_pdf(self, file_path: str) -> tuple[list[SectionNode], str]:
        """Parse a PDF file into section nodes and raw text."""
        all_text_parts: list[str] = []
        page_count = 0

        with pdfplumber.open(file_path) as pdf:
            page_count = len(pdf.pages)
            for page in pdf.pages:
                text = page.extract_text()
                if text:
                    all_text_parts.append(text)

        raw_text = "\n".join(all_text_parts)

        # Heuristic heading detection
        flat_sections: list[dict[str, Any]] = []
        current_content_parts: list[str] = []

        for line in raw_text.split("\n"):
            stripped = line.strip()
            if not stripped:
                continue
            # Heuristic: short line, all caps OR no ending period → heading
            if (
                len(stripped) < 80
                and (stripped.isupper() or not stripped.endswith("."))
                and len(stripped.split()) <= 12
                and len(stripped) > 0
            ):
                # Check if it looks like a heading (not just a short sentence)
                # Additional filter: must be relatively short and not contain common sentence patterns
                is_heading = stripped.isupper() or (
                    not stripped.endswith(".")
                    and not stripped.endswith(",")
                    and not stripped.endswith(";")
                    and not stripped.endswith(":")
                    and len(stripped) < 60
                )
                if is_heading:
                    if flat_sections:
                        flat_sections[-1]["content"] = "\n".join(
                            current_content_parts
                        ).strip()
                    current_content_parts = []
                    flat_sections.append({"title": stripped, "level": 1, "content": ""})
                    continue

            current_content_parts.append(stripped)

        if flat_sections:
            flat_sections[-1]["content"] = "\n".join(current_content_parts).strip()

        # If no headings found, wrap everything as single node
        if not flat_sections and raw_text.strip():
            flat_sections = [{"title": "全文", "level": 1, "content": raw_text.strip()}]

        tree = self._build_tree(flat_sections)
        return tree, raw_text

    # ------------------------------------------------------------------
    # Markdown parser
    # ------------------------------------------------------------------
    def _parse_md(self, file_path: str) -> tuple[list[SectionNode], str]:
        """Parse a Markdown file into section nodes and raw text."""
        with open(file_path, "r", encoding="utf-8") as f:
            content = f.read()

        raw_text = content
        flat_sections: list[dict[str, Any]] = []
        current_content_parts: list[str] = []

        for line in content.split("\n"):
            heading_match = re.match(r"^(#{1,6})\s+(.+)$", line)
            if heading_match:
                level = len(heading_match.group(1))
                title = heading_match.group(2).strip()
                # Flush content to previous section
                if flat_sections:
                    flat_sections[-1]["content"] = "\n".join(
                        current_content_parts
                    ).strip()
                current_content_parts = []
                flat_sections.append({"title": title, "level": level, "content": ""})
            else:
                stripped = line.strip()
                if stripped:
                    current_content_parts.append(stripped)

        # Flush last section's content
        if flat_sections:
            flat_sections[-1]["content"] = "\n".join(current_content_parts).strip()

        tree = self._build_tree(flat_sections)
        return tree, raw_text

    # ------------------------------------------------------------------
    # Tree builder — stack-based algorithm
    # ------------------------------------------------------------------
    def _build_tree(self, flat_sections: list[dict[str, Any]]) -> list[SectionNode]:
        """Convert a flat list of {title, level, content} into a nested tree.

        Uses a stack-based approach: each entry on the stack is (level, SectionNode).
        When a new section has a level > stack top, it becomes a child.
        When level <= stack top, pop until we find a suitable parent.
        """
        if not flat_sections:
            return []

        root: list[SectionNode] = []
        # Stack stores (level, node) pairs
        stack: list[tuple[int, SectionNode]] = []

        for section in flat_sections:
            node = SectionNode(
                title=section["title"],
                level=section["level"],
                content=section["content"],
            )

            # Pop stack until we find a parent with lower level
            while stack and stack[-1][0] >= node.level:
                stack.pop()

            if stack:
                # Attach as child of current stack top
                stack[-1][1].children.append(node)
            else:
                # Top-level node
                root.append(node)

            stack.append((node.level, node))

        return root


# ------------------------------------------------------------------
# Registration
# ------------------------------------------------------------------
from skills.registry import get_skill_registry  # noqa: E402

get_skill_registry().register(DocumentParseSkill())
