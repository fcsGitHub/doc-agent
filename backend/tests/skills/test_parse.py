"""Tests for DocumentParseSkill — DOCX, PDF, Markdown parsing."""

from __future__ import annotations

import os
import tempfile
from unittest.mock import MagicMock, patch

import pytest

from skills.base import SkillContext
from skills.parse import DocumentParseSkill, SectionNode


@pytest.fixture
def skill() -> DocumentParseSkill:
    """Fresh DocumentParseSkill instance."""
    return DocumentParseSkill()


@pytest.fixture
def ctx() -> SkillContext:
    """Minimal SkillContext for testing."""
    return SkillContext(task_id="test-task-001")


# ------------------------------------------------------------------
# Test 1: Parse DOCX with headings (H1 + 2×H2)
# ------------------------------------------------------------------
@pytest.mark.asyncio
async def test_parse_docx_with_headings(
    skill: DocumentParseSkill, ctx: SkillContext
) -> None:
    """Parse a DOCX with 1 H1 and 2 H2 headings. Expect 1 root with 2 children."""
    from docx import Document

    with tempfile.NamedTemporaryFile(suffix=".docx", delete=False) as tmp:
        tmp_path = tmp.name

    try:
        doc = Document()
        doc.add_heading("Main Title", level=1)
        doc.add_paragraph("Introduction paragraph.")
        doc.add_heading("Section A", level=2)
        doc.add_paragraph("Content of section A.")
        doc.add_heading("Section B", level=2)
        doc.add_paragraph("Content of section B.")
        doc.save(tmp_path)

        ctx.input_data = {"file_path": tmp_path, "file_type": "docx"}
        result = await skill.execute(ctx)

        assert result.success is True
        structure = result.output["structure"]
        assert len(structure) == 1  # 1 root (H1)
        root = structure[0]
        assert root["title"] == "Main Title"
        assert root["level"] == 1
        assert len(root["children"]) == 2
        assert root["children"][0]["title"] == "Section A"
        assert root["children"][1]["title"] == "Section B"

        # Raw text contains paragraph content
        raw_text = result.output["raw_text"]
        assert "Introduction paragraph" in raw_text
        assert "Content of section A" in raw_text
        assert "Content of section B" in raw_text
    finally:
        os.unlink(tmp_path)


# ------------------------------------------------------------------
# Test 2: Parse Markdown with nested headings
# ------------------------------------------------------------------
@pytest.mark.asyncio
async def test_parse_markdown_nested(
    skill: DocumentParseSkill, ctx: SkillContext
) -> None:
    """Parse Markdown with H1 and H2 headings; verify tree structure."""
    md_content = "# Title\n\nPara1\n\n## Section A\n\nPara2\n\n## Section B\n\nPara3\n"

    with tempfile.NamedTemporaryFile(
        suffix=".md", mode="w", delete=False, encoding="utf-8"
    ) as tmp:
        tmp.write(md_content)
        tmp_path = tmp.name

    try:
        ctx.input_data = {"file_path": tmp_path, "file_type": "md"}
        result = await skill.execute(ctx)

        assert result.success is True
        structure = result.output["structure"]
        assert len(structure) == 1  # 1 root (H1)

        root = structure[0]
        assert root["title"] == "Title"
        assert root["level"] == 1
        assert root["content"] == "Para1"
        assert len(root["children"]) == 2
        assert root["children"][0]["title"] == "Section A"
        assert root["children"][0]["content"] == "Para2"
        assert root["children"][1]["title"] == "Section B"
        assert root["children"][1]["content"] == "Para3"
    finally:
        os.unlink(tmp_path)


# ------------------------------------------------------------------
# Test 3: Empty DOCX handled gracefully
# ------------------------------------------------------------------
@pytest.mark.asyncio
async def test_parse_empty_docx(skill: DocumentParseSkill, ctx: SkillContext) -> None:
    """Empty DOCX returns empty structure and empty raw_text."""
    from docx import Document

    with tempfile.NamedTemporaryFile(suffix=".docx", delete=False) as tmp:
        tmp_path = tmp.name

    try:
        doc = Document()
        doc.save(tmp_path)

        ctx.input_data = {"file_path": tmp_path, "file_type": "docx"}
        result = await skill.execute(ctx)

        assert result.success is True
        assert result.output["structure"] == []
        assert result.output["raw_text"] == ""
        assert result.output["metadata"]["word_count"] == 0
    finally:
        os.unlink(tmp_path)


# ------------------------------------------------------------------
# Test 4: Tree building from flat list
# ------------------------------------------------------------------
def test_build_tree_directly(skill: DocumentParseSkill) -> None:
    """Test _build_tree with known flat input produces correct nesting."""
    flat = [
        {"title": "Chapter 1", "level": 1, "content": "Intro"},
        {"title": "Section 1.1", "level": 2, "content": "Details A"},
        {"title": "Section 1.2", "level": 2, "content": "Details B"},
        {"title": "Chapter 2", "level": 1, "content": "More"},
        {"title": "Section 2.1", "level": 2, "content": "Details C"},
        {"title": "Sub 2.1.1", "level": 3, "content": "Deep"},
    ]
    tree = skill._build_tree(flat)

    # 2 top-level nodes
    assert len(tree) == 2
    assert tree[0].title == "Chapter 1"
    assert len(tree[0].children) == 2
    assert tree[0].children[0].title == "Section 1.1"
    assert tree[0].children[1].title == "Section 1.2"

    assert tree[1].title == "Chapter 2"
    assert len(tree[1].children) == 1
    assert tree[1].children[0].title == "Section 2.1"
    assert len(tree[1].children[0].children) == 1
    assert tree[1].children[0].children[0].title == "Sub 2.1.1"


def test_build_tree_empty(skill: DocumentParseSkill) -> None:
    """Empty flat list returns empty tree."""
    assert skill._build_tree([]) == []


# ------------------------------------------------------------------
# Test 5: PDF parsing with mocked pdfplumber
# ------------------------------------------------------------------
@pytest.mark.asyncio
async def test_parse_pdf_mocked(skill: DocumentParseSkill, ctx: SkillContext) -> None:
    """PDF parsing with mocked pdfplumber pages."""
    # Create mock page objects
    page1 = MagicMock()
    page1.extract_text.return_value = (
        "INTRODUCTION\nThis is the first page content.\nMore text here."
    )

    page2 = MagicMock()
    page2.extract_text.return_value = (
        "METHODOLOGY\nWe used various approaches.\nResults follow."
    )

    mock_pdf = MagicMock()
    mock_pdf.pages = [page1, page2]
    mock_pdf.__enter__ = MagicMock(return_value=mock_pdf)
    mock_pdf.__exit__ = MagicMock(return_value=False)

    with patch("skills.parse.pdfplumber") as mock_pdfplumber:
        mock_pdfplumber.open.return_value = mock_pdf

        # Need a dummy file path — won't actually be opened
        ctx.input_data = {"file_path": "/fake/test.pdf", "file_type": "pdf"}
        result = await skill.execute(ctx)

    assert result.success is True
    raw_text = result.output["raw_text"]
    assert "INTRODUCTION" in raw_text
    assert "METHODOLOGY" in raw_text

    structure = result.output["structure"]
    # Should detect INTRODUCTION and METHODOLOGY as headings
    assert len(structure) >= 1
    titles = [s["title"] for s in structure]
    assert "INTRODUCTION" in titles
    assert "METHODOLOGY" in titles


# ------------------------------------------------------------------
# Test 6: Skill registration in global registry
# ------------------------------------------------------------------
def test_skill_registration() -> None:
    """DocumentParseSkill is registered in the global skill registry."""
    from skills.registry import get_skill_registry

    registry = get_skill_registry()
    assert registry.is_registered("document_parse")
    skill = registry.get("document_parse")
    assert isinstance(skill, DocumentParseSkill)


# ------------------------------------------------------------------
# Test 7: Unsupported file type returns failure
# ------------------------------------------------------------------
@pytest.mark.asyncio
async def test_unsupported_file_type(
    skill: DocumentParseSkill, ctx: SkillContext
) -> None:
    """Unsupported file type returns SkillResult with success=False."""
    ctx.input_data = {"file_path": "/fake/file.txt", "file_type": "txt"}
    result = await skill.execute(ctx)

    assert result.success is False
    assert "Unsupported file type" in (result.error or "")


# ------------------------------------------------------------------
# Test 8: Markdown with only content (no headings)
# ------------------------------------------------------------------
@pytest.mark.asyncio
async def test_parse_markdown_no_headings(
    skill: DocumentParseSkill, ctx: SkillContext
) -> None:
    """Markdown without headings returns empty structure."""
    md_content = "Just a paragraph.\nAnother paragraph.\n"

    with tempfile.NamedTemporaryFile(
        suffix=".md", mode="w", delete=False, encoding="utf-8"
    ) as tmp:
        tmp.write(md_content)
        tmp_path = tmp.name

    try:
        ctx.input_data = {"file_path": tmp_path, "file_type": "md"}
        result = await skill.execute(ctx)

        assert result.success is True
        assert result.output["structure"] == []
        assert "Just a paragraph" in result.output["raw_text"]
    finally:
        os.unlink(tmp_path)
