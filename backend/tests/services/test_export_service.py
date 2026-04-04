"""Unit tests for ExportService — mocked async DB session."""

from __future__ import annotations

import uuid
from io import BytesIO
from typing import Any
from unittest.mock import AsyncMock, MagicMock

import pytest
from docx import Document

from services.export_service import ExportService


# ---------------------------------------------------------------------------
# Helpers
# ---------------------------------------------------------------------------


def _make_task(**overrides: Any) -> MagicMock:
    """Return a mock Task ORM object."""
    defaults: dict[str, Any] = {
        "id": uuid.uuid4(),
        "name": "Test Report",
        "doc_type": "report",
        "status": "completed",
    }
    defaults.update(overrides)
    task = MagicMock(**defaults)
    for k, v in defaults.items():
        setattr(task, k, v)
    return task


def _make_section(**overrides: Any) -> MagicMock:
    """Return a mock Section ORM object."""
    defaults: dict[str, Any] = {
        "id": uuid.uuid4(),
        "task_id": uuid.uuid4(),
        "title": "Introduction",
        "level": 1,
        "order_index": 0,
        "status": "generated",
    }
    defaults.update(overrides)
    section = MagicMock(**defaults)
    for k, v in defaults.items():
        setattr(section, k, v)
    return section


def _make_version(**overrides: Any) -> MagicMock:
    """Return a mock SectionVersion ORM object."""
    defaults: dict[str, Any] = {
        "id": uuid.uuid4(),
        "section_id": uuid.uuid4(),
        "version_number": 1,
        "content": "Some content here.",
        "word_count": 3,
    }
    defaults.update(overrides)
    version = MagicMock(**defaults)
    for k, v in defaults.items():
        setattr(version, k, v)
    return version


def _mock_db() -> AsyncMock:
    """Create a fully mocked AsyncSession."""
    db = AsyncMock()
    db.add = MagicMock()
    db.flush = AsyncMock()
    db.commit = AsyncMock()
    db.refresh = AsyncMock()
    db.execute = AsyncMock()
    return db


def _scalar_result(value: Any) -> MagicMock:
    """Mock db.execute() result with scalar_one_or_none returning value."""
    result = MagicMock()
    result.scalar_one_or_none.return_value = value
    return result


def _scalars_result(values: list[Any]) -> MagicMock:
    """Mock db.execute() result with scalars().all() returning values."""
    result = MagicMock()
    scalars = MagicMock()
    scalars.all.return_value = values
    result.scalars.return_value = scalars
    return result


def _read_docx(docx_bytes: bytes) -> Document:
    """Parse DOCX bytes into a python-docx Document for inspection."""
    return Document(BytesIO(docx_bytes))


# ---------------------------------------------------------------------------
# Tests
# ---------------------------------------------------------------------------


class TestExportDocxHeadings:
    """Test DOCX created with correct headings."""

    @pytest.mark.asyncio
    async def test_task_title_as_h1_and_sections_as_h2(self) -> None:
        """Task title should be Heading 1, section titles should be Heading 2."""
        db = _mock_db()
        task_id = uuid.uuid4()
        sec1_id = uuid.uuid4()
        sec2_id = uuid.uuid4()

        task_mock = _make_task(id=task_id, name="Annual Report 2026")
        sec1 = _make_section(
            id=sec1_id, task_id=task_id, title="Executive Summary", order_index=0
        )
        sec2 = _make_section(
            id=sec2_id, task_id=task_id, title="Financial Overview", order_index=1
        )

        v1 = _make_version(
            section_id=sec1_id, version_number=1, content="Summary content here."
        )
        v2 = _make_version(
            section_id=sec2_id, version_number=1, content="Financial details."
        )

        # Call 1: task lookup, Call 2: sections, Call 3: versions
        db.execute.side_effect = [
            _scalar_result(task_mock),
            _scalars_result([sec1, sec2]),
            _scalars_result([v1, v2]),
        ]

        svc = ExportService()
        docx_bytes, task_name = await svc.export_docx(db, str(task_id))

        assert task_name == "Annual Report 2026"
        assert len(docx_bytes) > 0

        doc = _read_docx(docx_bytes)
        paragraphs = doc.paragraphs

        # Find headings
        headings = [
            (p.text, p.style.name)
            for p in paragraphs
            if "Heading" in (p.style.name or "")
        ]

        assert ("Annual Report 2026", "Heading 1") in headings
        assert ("Executive Summary", "Heading 2") in headings
        assert ("Financial Overview", "Heading 2") in headings

    @pytest.mark.asyncio
    async def test_latest_version_used(self) -> None:
        """When multiple versions exist, the highest version_number is used."""
        db = _mock_db()
        task_id = uuid.uuid4()
        sec_id = uuid.uuid4()

        task_mock = _make_task(id=task_id, name="Versioned Doc")
        sec = _make_section(
            id=sec_id, task_id=task_id, title="Chapter 1", order_index=0
        )

        v1 = _make_version(section_id=sec_id, version_number=1, content="Old content.")
        v2 = _make_version(section_id=sec_id, version_number=2, content="New content.")
        v3 = _make_version(
            section_id=sec_id, version_number=3, content="Latest content."
        )

        db.execute.side_effect = [
            _scalar_result(task_mock),
            _scalars_result([sec]),
            _scalars_result([v1, v2, v3]),
        ]

        svc = ExportService()
        docx_bytes, _ = await svc.export_docx(db, str(task_id))

        doc = _read_docx(docx_bytes)
        all_text = " ".join(p.text for p in doc.paragraphs)

        assert "Latest content." in all_text
        assert "Old content." not in all_text


class TestExportDocxMarkdown:
    """Test markdown formatting in DOCX output."""

    @pytest.mark.asyncio
    async def test_markdown_h3_bold_italic_bullet(self) -> None:
        """## headings -> H3, **bold** -> bold run, *italic* -> italic run, - item -> bullet."""
        db = _mock_db()
        task_id = uuid.uuid4()
        sec_id = uuid.uuid4()

        task_mock = _make_task(id=task_id, name="Markdown Test")
        sec = _make_section(
            id=sec_id, task_id=task_id, title="Section A", order_index=0
        )

        markdown_content = (
            "## Subsection Title\n"
            "This has **bold text** and *italic text* in it.\n"
            "- First bullet item\n"
            "- Second bullet item"
        )
        v1 = _make_version(
            section_id=sec_id, version_number=1, content=markdown_content
        )

        db.execute.side_effect = [
            _scalar_result(task_mock),
            _scalars_result([sec]),
            _scalars_result([v1]),
        ]

        svc = ExportService()
        docx_bytes, _ = await svc.export_docx(db, str(task_id))

        doc = _read_docx(docx_bytes)

        # Check ## -> Heading 3
        h3_paras = [p for p in doc.paragraphs if p.style.name == "Heading 3"]
        assert len(h3_paras) == 1
        assert h3_paras[0].text == "Subsection Title"

        # Check bold run exists
        bold_runs = []
        italic_runs = []
        for p in doc.paragraphs:
            for run in p.runs:
                if run.bold:
                    bold_runs.append(run.text)
                if run.italic:
                    italic_runs.append(run.text)

        assert "bold text" in bold_runs
        assert "italic text" in italic_runs

        # Check bullet paragraphs
        bullet_paras = [p for p in doc.paragraphs if p.style.name == "List Bullet"]
        assert len(bullet_paras) == 2
        bullet_texts = [p.text for p in bullet_paras]
        assert "First bullet item" in bullet_texts
        assert "Second bullet item" in bullet_texts


class TestExportDocx404:
    """Test 404 when task not found."""

    @pytest.mark.asyncio
    async def test_task_not_found_raises_404(self) -> None:
        """export_docx raises HTTPException 404 when task doesn't exist."""
        db = _mock_db()
        db.execute.return_value = _scalar_result(None)

        svc = ExportService()
        from fastapi import HTTPException

        with pytest.raises(HTTPException) as exc_info:
            await svc.export_docx(db, str(uuid.uuid4()))

        assert exc_info.value.status_code == 404
        assert "Task not found" in str(exc_info.value.detail)


class TestExportDocxEmptySections:
    """Test DOCX generation when task has no sections."""

    @pytest.mark.asyncio
    async def test_empty_sections_produces_valid_docx(self) -> None:
        """A task with no sections should still produce a valid DOCX with just the title."""
        db = _mock_db()
        task_id = uuid.uuid4()

        task_mock = _make_task(id=task_id, name="Empty Task")

        db.execute.side_effect = [
            _scalar_result(task_mock),
            _scalars_result([]),  # No sections
        ]

        svc = ExportService()
        docx_bytes, task_name = await svc.export_docx(db, str(task_id))

        assert task_name == "Empty Task"
        assert len(docx_bytes) > 0

        doc = _read_docx(docx_bytes)
        headings = [p for p in doc.paragraphs if p.style.name == "Heading 1"]
        assert len(headings) == 1
        assert headings[0].text == "Empty Task"
