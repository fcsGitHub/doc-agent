"""Export service — DOCX generation from task sections."""

from __future__ import annotations

import re
import uuid
from collections import defaultdict
from io import BytesIO
from typing import Any

from docx import Document  # pyright: ignore[reportGeneralTypeIssues]
from fastapi import HTTPException
from sqlalchemy import select
from sqlalchemy.ext.asyncio import AsyncSession

from models.section import Section, SectionVersion
from models.task import Task
from services.audit_service import AuditService


class ExportService:
    """Service for exporting task documents to DOCX format."""

    _audit = AuditService()

    async def export_docx(self, db: AsyncSession, task_id: str) -> tuple[bytes, str]:
        """Generate a DOCX file from task sections.

        Returns:
            Tuple of (docx_bytes, task_name) for the response.

        Raises:
            HTTPException: 404 if task not found.
        """
        # Load task
        stmt = select(Task).where(Task.id == uuid.UUID(task_id))
        result = await db.execute(stmt)
        task = result.scalar_one_or_none()

        if task is None:
            raise HTTPException(status_code=404, detail="Task not found")

        # Load all sections for the task, ordered by order_index
        sec_stmt = (
            select(Section)
            .where(Section.task_id == uuid.UUID(task_id))
            .order_by(Section.order_index)
        )
        sec_result = await db.execute(sec_stmt)
        sections = list(sec_result.scalars().all())

        # Load all section versions and find latest per section
        latest_content: dict[uuid.UUID, str] = {}
        if sections:
            sec_ids = [s.id for s in sections]
            v_stmt = select(SectionVersion).where(
                SectionVersion.section_id.in_(sec_ids)
            )
            v_result = await db.execute(v_stmt)
            all_versions = list(v_result.scalars().all())

            # Group by section_id, pick highest version_number
            versions_by_sec: dict[uuid.UUID, list[SectionVersion]] = defaultdict(list)
            for v in all_versions:
                versions_by_sec[v.section_id].append(v)

            for sec_id, versions in versions_by_sec.items():
                latest = max(versions, key=lambda v: v.version_number)
                latest_content[sec_id] = latest.content

        # Build DOCX
        doc = Document()

        # Task title as Heading 1
        doc.add_heading(task.name, level=1)

        # Each section
        for section in sections:
            # Section title as Heading 2
            doc.add_heading(section.title, level=2)

            # Get content for this section
            content = latest_content.get(section.id, "")
            if content:
                _render_markdown_content(doc, content)

        # Save to bytes
        buffer = BytesIO()
        doc.save(buffer)
        buffer.seek(0)

        # Best-effort audit logging
        try:
            await self._audit.log(
                db,
                task_id=task_id,
                action="document_exported",
                entity_type="task",
                entity_id=task_id,
                details={"format": "docx", "task_name": task.name},
            )
        except Exception:
            pass

        return buffer.getvalue(), task.name


def _render_markdown_content(doc: Any, content: str) -> None:
    """Render markdown content into the DOCX document.

    Supports:
    - ## headings -> Heading 3
    - **bold** -> bold runs
    - *italic* -> italic runs
    - - list items -> List Bullet style
    - Regular paragraphs -> Normal style
    """
    lines = content.split("\n")
    i = 0
    while i < len(lines):
        line = lines[i].rstrip()

        # Skip empty lines
        if not line.strip():
            i += 1
            continue

        # ## Subsection headings -> Heading 3
        heading_match = re.match(r"^##\s+(.+)$", line)
        if heading_match:
            doc.add_heading(heading_match.group(1).strip(), level=3)
            i += 1
            continue

        # Bullet list items: - item
        if re.match(r"^[-*]\s+", line):
            text = re.sub(r"^[-*]\s+", "", line)
            para = doc.add_paragraph(style="List Bullet")
            _add_formatted_runs(para, text)
            i += 1
            continue

        # Regular paragraph
        para = doc.add_paragraph()
        _add_formatted_runs(para, line)
        i += 1


def _add_formatted_runs(paragraph, text: str) -> None:
    """Add runs to a paragraph, handling **bold** and *italic* markdown.

    Processes text left-to-right, splitting on bold/italic markers.
    """
    # Pattern matches **bold** or *italic* segments
    # Order matters: check ** before *
    pattern = re.compile(r"(\*\*(.+?)\*\*|\*(.+?)\*)")

    pos = 0
    for match in pattern.finditer(text):
        # Add text before the match as a normal run
        if match.start() > pos:
            paragraph.add_run(text[pos : match.start()])

        if match.group(2) is not None:
            # **bold**
            run = paragraph.add_run(match.group(2))
            run.bold = True
        elif match.group(3) is not None:
            # *italic*
            run = paragraph.add_run(match.group(3))
            run.italic = True

        pos = match.end()

    # Add remaining text after last match
    if pos < len(text):
        paragraph.add_run(text[pos:])
